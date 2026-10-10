"""Build the reviewable DOCX from the maintained Markdown development guide."""
import argparse
import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


def font(style, size, bold=False):
    style.font.name = "Noto Sans CJK SC"
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = RGBColor(0, 0, 0)
    rpr = style.element.get_or_add_rPr()
    for key in list(rpr.rFonts.attrib):
        if "theme" in key.lower():
            del rpr.rFonts.attrib[key]
    rpr.rFonts.set(qn("w:eastAsia"), "Noto Sans CJK SC")
    rpr.rFonts.set(qn("w:cs"), "Noto Sans CJK SC")
    lang = OxmlElement("w:lang")
    lang.set(qn("w:val"), "zh-CN")
    lang.set(qn("w:eastAsia"), "zh-CN")
    rpr.append(lang)


def plain(text):
    return re.sub(r"\[([^\]]+)\]\([^)]+\)", r"\1", text).replace("**", "").replace("`", "")


def table(doc, lines):
    rows = [[plain(c.strip()) for c in line.strip().strip("|").split("|")] for line in lines]
    rows = [r for r in rows if not all(re.fullmatch(r"[-: ]+", c) for c in r)]
    n = len(rows[0])
    widths = {2: [5.3, 11.9], 3: [3.2, 6.4, 7.6], 4: [2.7, 4.1, 5.2, 5.2]}[n]
    t = doc.add_table(rows=0, cols=n)
    t.autofit = False
    for c, w in zip(t.columns, widths):
        c.width = Cm(w)
    borders = OxmlElement("w:tblBorders")
    for side in ["top", "left", "bottom", "right", "insideH", "insideV"]:
        edge = OxmlElement(f"w:{side}")
        for key, value in [("val", "single"), ("sz", "4"), ("color", "D9D9D9")]:
            edge.set(qn(f"w:{key}"), value)
        borders.append(edge)
    t._tbl.tblPr.append(borders)
    for i, values in enumerate(rows):
        row = t.add_row()
        row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
        if i == 0:
            row._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
        for j, value in enumerate(values):
            cell = row.cells[j]
            cell.width = Cm(widths[j])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            pr = cell._tc.get_or_add_tcPr()
            margins = OxmlElement("w:tcMar")
            for side in ["top", "left", "bottom", "right"]:
                el = OxmlElement(f"w:{side}")
                el.set(qn("w:w"), "95")
                el.set(qn("w:type"), "dxa")
                margins.append(el)
            pr.append(margins)
            fill = OxmlElement("w:shd")
            fill.set(qn("w:fill"), "E7EEF5" if i == 0 else "FFFFFF")
            pr.append(fill)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.line_spacing = 1.15
            run = p.add_run(value)
            run.font.size = Pt(9.5)
            run.font.bold = i == 0
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def build(source, target):
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.top_margin, sec.bottom_margin = Cm(1.7), Cm(1.7)
    sec.left_margin, sec.right_margin = Cm(1.9), Cm(1.9)
    for name, size, bold in [("Normal", 10.5, False), ("Title", 22, True),
                             ("Heading 1", 16, True), ("Heading 2", 12, True),
                             ("Heading 3", 11, True), ("List Bullet", 10.5, False)]:
        font(doc.styles[name], size, bold)
    normal = doc.styles["Normal"].paragraph_format
    normal.space_after = Pt(6)
    normal.line_spacing = 1.2
    for h in ["Heading 1", "Heading 2", "Heading 3"]:
        pf = doc.styles[h].paragraph_format
        pf.space_before, pf.space_after = Pt(12), Pt(6)
        pf.keep_with_next = True
    for style in doc.styles:
        for border in list(style.element.iter(qn("w:pBdr"))):
            border.getparent().remove(border)
    for fonts in doc.styles.element.iter(qn("w:rFonts")):
        for key in list(fonts.attrib):
            if "theme" in key.lower():
                del fonts.attrib[key]
        for key in ["ascii", "hAnsi", "eastAsia", "cs"]:
            fonts.set(qn(f"w:{key}"), "Noto Sans CJK SC")
    footer = sec.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = footer.add_run()
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    r._r.addnext(field)
    pending = []
    for line in source.read_text(encoding="utf-8").splitlines():
        if line.startswith("|"):
            pending.append(line)
            continue
        if pending:
            table(doc, pending)
            pending = []
        if not line.strip():
            continue
        if line.startswith("# "):
            doc.add_paragraph(plain(line[2:]), "Title")
        elif line.startswith("## "):
            doc.add_paragraph(plain(line[3:]), "Heading 1")
        elif line.startswith("### "):
            doc.add_paragraph(plain(line[4:]), "Heading 2")
        elif line.startswith("- "):
            doc.add_paragraph(plain(line[2:]), "List Bullet")
        else:
            doc.add_paragraph(plain(line))
    if pending:
        table(doc, pending)
    doc.core_properties.title = "ClipForge 内容生产与开发指导"
    doc.core_properties.subject = "运营计划 内容生产能力 需求契约 验收标准"
    doc.core_properties.author = ""
    target.parent.mkdir(parents=True, exist_ok=True)
    doc.save(target)
    print(target.resolve())


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("source", type=Path)
    parser.add_argument("target", type=Path)
    args = parser.parse_args()
    build(args.source, args.target)
