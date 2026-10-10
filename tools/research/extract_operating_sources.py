"""Read supplied PDF/Office files without modifying them; retain source locators."""
import argparse
import hashlib
import json
from pathlib import Path
from zipfile import ZipFile

from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph
from openpyxl import load_workbook
from pypdf import PdfReader


def extract(path):
    result = {"name": path.name, "path": str(path),
              "sha256": hashlib.sha256(path.read_bytes()).hexdigest()}
    lines = [f"# {path.name}"]
    if path.suffix.lower() == ".pdf":
        reader = PdfReader(path)
        result["pages"] = len(reader.pages)
        for i, page in enumerate(reader.pages, 1):
            lines.extend([f"\n## Page {i}", page.extract_text() or "[No extractable text]"])
    elif path.suffix.lower() == ".xlsx":
        wb = load_workbook(path, read_only=True, data_only=False)
        result["sheets"] = []
        for ws in wb:
            result["sheets"].append({"name": ws.title, "rows": ws.max_row, "columns": ws.max_column})
            lines.append(f"\n## Sheet {ws.title} ({ws.max_row} rows x {ws.max_column} columns)")
            for row in ws.iter_rows():
                values = [f"{c.coordinate}={json.dumps(c.value, ensure_ascii=False, default=str)}"
                          for c in row if c.value is not None]
                if values:
                    lines.append(" | ".join(values))
        wb.close()
        with ZipFile(path) as z:
            result["embedded_media"] = [n for n in z.namelist() if n.startswith("xl/media/")]
    elif path.suffix.lower() == ".docx":
        doc = Document(path)
        for i, block in enumerate(doc.element.body, 1):
            if block.tag.endswith("}p"):
                p = Paragraph(block, doc)
                if p.text.strip():
                    lines.append(f"B{i} [{p.style.name}] {p.text}")
            elif block.tag.endswith("}tbl"):
                lines.append(f"\n### Table B{i}")
                for r, row in enumerate(Table(block, doc).rows, 1):
                    lines.append(f"R{r}: " + " | ".join(c.text for c in row.cells))
        with ZipFile(path) as z:
            result["embedded_media"] = [n for n in z.namelist() if n.startswith("word/media/")]
    else:
        raise ValueError(f"Unsupported source: {path}")
    return result, "\n".join(lines)


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--out", type=Path, required=True)
    parser.add_argument("inputs", nargs="+", type=Path)
    args = parser.parse_args()
    args.out.mkdir(parents=True, exist_ok=True)
    manifest = []
    for path in args.inputs:
        item, body = extract(path.resolve(strict=True))
        (args.out / (path.stem + ".txt")).write_text(body, encoding="utf-8")
        manifest.append(item)
    (args.out / "source-manifest.json").write_text(
        json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(manifest, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
