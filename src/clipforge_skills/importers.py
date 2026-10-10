"""Preserve legacy spreadsheet and asset-hub evidence without trusting examples."""
import json
from pathlib import Path
from zipfile import ZipFile

from openpyxl import load_workbook

from .core import Fault, digest, write_json
from .library import Library


def _read_sheet(sheet):
    if sheet.max_row * sheet.max_column > 1000000:
        raise Fault("workbook_too_large", sheet.title)
    cells = []
    for row in sheet:
        for cell in row:
            if cell.value is None:
                continue
            value = cell.value
            if not isinstance(value, (str, int, float, bool)):
                value = str(value)
            cells.append({"cell": cell.coordinate, "value": value, "type": cell.data_type})
    return {"name": sheet.title, "cells": cells}


def _index_rows(sheet):
    rows = list(sheet.iter_rows(values_only=True))
    for i, header in enumerate(rows[:5]):
        if "文件名" not in header or "存储路径" not in header:
            continue
        result = []
        for row_no, values in enumerate(rows[i + 1:], i + 2):
            record = {str(k): v for k, v in zip(header, values) if k and v is not None}
            if record.get("文件名"):
                result.append({"sheet": sheet.title, "row": row_no, "fields": record,
                               "media_status": "not_resolved", "facts_status": "pending"})
        return result
    return []


def _read_images(path, out):
    images = []
    with ZipFile(path) as archive:
        for member in archive.infolist():
            if not member.filename.startswith("xl/media/") or member.is_dir():
                continue
            if member.file_size > 20 * 1024 * 1024:
                raise Fault("image_too_large", member.filename)
            dest = out / "reference-images" / Path(member.filename).name
            dest.parent.mkdir(parents=True, exist_ok=True)
            dest.write_bytes(archive.read(member))
            images.append({"source_member": member.filename, "file": str(dest),
                           "purpose": "reference_only", "sha256": digest(dest)})
        anchors = {n: archive.read(n).decode("utf-8") for n in archive.namelist()
                   if (n.startswith("xl/drawings/") or n.startswith("xl/worksheets/_rels/"))
                   and n.endswith((".xml", ".rels"))}
    return images, anchors


def import_docs(root, paths):
    summary = []
    for name in paths:
        path = Path(name).resolve()
        if path.suffix.lower() != ".xlsx" or not path.is_file():
            raise Fault("invalid_workbook", str(path))
        key = digest(path)
        out = Path(root) / "imports" / key
        workbook = load_workbook(path, read_only=True, data_only=False)
        try:
            sheets = [_read_sheet(sheet) for sheet in workbook]
            index_rows = [row for sheet in workbook for row in _index_rows(sheet)]
        finally:
            workbook.close()
        images, anchors = _read_images(path, out)
        manifest = {"source": str(path), "sha256": key, "facts_status": "pending", "sheets": sheets,
                    "images": images, "drawing_anchors": anchors, "index_rows": index_rows}
        write_json(out / "source.json", manifest)
        summary.append({"source": path.name, "worksheets": len(sheets), "images": len(images),
                        "unresolved_media_rows": len(index_rows), "artifact": str(out / "source.json")})
    return {"sources": summary, "worksheets": sum(x["worksheets"] for x in summary),
            "images": sum(x["images"] for x in summary)}


def _legacy_rows(index):
    path = Path(index)
    if path.stat().st_size > 8 * 1024 * 1024:
        raise Fault("invalid_hub", "Split large legacy indices before import")
    data = json.loads(path.read_text(encoding="utf-8"))
    version = data.get("schema_version", 1) if isinstance(data, dict) else 1
    if str(version).split(".")[0] != "1":
        raise Fault("unsupported_hub_schema", str(version))
    rows = data.get("rows", data.get("shots", data.get("items", []))) if isinstance(data, dict) else data
    if not isinstance(rows, list) or any(not isinstance(row, dict) for row in rows):
        raise Fault("invalid_hub", "Expected an array of objects")
    return rows


def _legacy_labels(row, asset, index):
    aid = asset["asset_id"]
    start = float(row.get("start_time", 0))
    end = float(row.get("end_time", asset["technical"]["duration"]))
    if not 0 <= start < end <= asset["technical"]["duration"] + .001:
        raise Fault("invalid_hub_time", str(row.get("编号", aid)))
    return [{"label": key, "value": str(row[key]), "start": start, "end": end,
             "source": "legacy", "evidence": f"{Path(index).resolve()}#{row.get('编号', aid)}",
             "version": "legacy-v1", "review_status": "pending"}
            for key in ("语义描述", "行为", "景别", "情绪", "二级分类建议") if row.get(key)]


def _import_legacy_file(root, library, path, row, index):
    # Register precisely this file without copying it or requiring symlink rights.
    result = library.ingest_file(path)
    for aid in result["asset_ids"]:
        library.annotate(aid, {"annotations": _legacy_labels(row, library.get(aid), index)})
    return {"original": row, "asset_ids": result["asset_ids"], "errors": result["errors"],
            "labels_status": "pending"}


def import_hub(root, index, media_root):
    mapped, missing = [], []
    library = Library(root)
    for row in _legacy_rows(index):
        raw = row.get("file_path", row.get("path", row.get("存储路径")))
        if not raw:
            missing.append({"original": row, "reason": "missing_path"})
            continue
        path = Path(raw)
        path = path if path.is_absolute() else Path(media_root) / path
        if not path.is_file():
            missing.append({"original": row, "reason": "missing_file"})
            continue
        mapped.append(_import_legacy_file(root, library, path, row, index))
    out = Path(root) / "imports" / digest(index) / "hub.json"
    write_json(out, {"source": str(Path(index).resolve()), "mapped": mapped, "missing": missing})
    return {"mapped": len(mapped), "missing": len(missing), "artifact": str(out)}
